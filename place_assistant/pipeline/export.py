"""Writes collected place records and product surveys to disk in three formats:

- JSON: full fidelity.
- CSV: flattened nested fields, for spreadsheets.
- DOCX: a readable per-place report. It is a local file that opens in Google
  Docs (File > Open, or drag it into Drive). No Google Docs API is involved.
"""

import csv
import json
import os
import re

from docx import Document
from docx.shared import Pt

from . import verdict


def safe_file_stem(*parts: str) -> str:
    """Joins parts into a file name stem that can't contain a folder, e.g.
    ("Menteng / Jakarta", "cafe") -> "menteng___jakarta_cafe".
    """
    stem = "_".join(p for p in parts if p).lower().replace(",", "").replace(" ", "_")
    return re.sub(r"[^\w.\-]", "_", stem).strip("._") or "output"


def _remove_stale(path: str) -> None:
    """Deletes a CSV or DOCX left by an earlier run with the same name, so it
    never sits next to a newer, empty JSON as if it belonged to it.
    """
    try:
        os.remove(path)
    except FileNotFoundError:
        pass


def save_records(records: list[dict], output_dir: str | os.PathLike, base_name: str) -> dict:
    """Writes JSON always and CSV only when there are records. Returns both
    paths. An older CSV of the same name is removed when there are none.
    """
    os.makedirs(output_dir, exist_ok=True)
    json_path = os.path.join(output_dir, f"{base_name}.json")
    csv_path = os.path.join(output_dir, f"{base_name}.csv")

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    if records:
        flat_rows = [_flatten(r) for r in records]
        fieldnames = sorted({key for row in flat_rows for key in row})
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(flat_rows)
    else:
        _remove_stale(csv_path)

    return {"json_path": json_path, "csv_path": csv_path if records else None}


def save_docx(records: list[dict], output_dir: str | os.PathLike, base_name: str, title: str) -> str | None:
    """Writes a readable per-place report (address, rating, facilities,
    sentiment, and a menu/price table) to a .docx file. Returns None (writes
    nothing, and removes an older .docx of the same name) if records is
    empty — there's nothing meaningful to report.
    """
    if not records:
        _remove_stale(os.path.join(output_dir, f"{base_name}.docx"))
        return None

    os.makedirs(output_dir, exist_ok=True)
    docx_path = os.path.join(output_dir, f"{base_name}.docx")

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"{len(records)} place(s) collected.")

    for record in records:
        doc.add_heading(record.get("name") or "Unknown", level=1)

        facts = doc.add_paragraph()
        facts.add_run("Address: ").bold = True
        facts.add_run(record.get("address") or "not available")
        facts.add_run("\n")
        facts.add_run("Rating: ").bold = True
        rating = record.get("rating")
        total_reviews = record.get("total_reviews")
        facts.add_run(
            f"{rating} ({total_reviews} reviews)" if rating is not None else "not available"
        )
        facts.add_run("\n")
        facts.add_run("Price level: ").bold = True
        facts.add_run(str(record.get("price_level") or "not available"))
        facts.add_run("\n")
        facts.add_run("Open now: ").bold = True
        facts.add_run(str(record.get("open_now")))
        facts.add_run("\n")
        facts.add_run("Google Maps: ").bold = True
        facts.add_run(record.get("google_maps_url") or "not available")

        alcohol = record.get("alcohol") or {}
        if alcohol.get("serves_alcohol"):
            doc.add_paragraph(
                "Serves alcohol: "
                + ", ".join(
                    name
                    for name, flag in [
                        ("beer", alcohol.get("serves_beer")),
                        ("wine", alcohol.get("serves_wine")),
                        ("cocktails", alcohol.get("serves_cocktails")),
                    ]
                    if flag
                )
            )

        doc.add_heading("Menu & Prices", level=2)
        menu = record.get("menu") or {}
        items = menu.get("items") or []
        source = menu.get("source", "unavailable")
        if items:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            header_cells = table.rows[0].cells
            header_cells[0].text = "Item"
            header_cells[1].text = "Price"
            for entry in items:
                row_cells = table.add_row().cells
                row_cells[0].text = str(entry.get("item", ""))
                row_cells[1].text = str(entry.get("price", ""))
            source_p = doc.add_paragraph()
            source_p.add_run(
                "Source: official website"
                if source == "official_website"
                else "Source: delivery-platform/web search — prices may be outdated, not from the business's own site"
            ).italic = True
        else:
            doc.add_paragraph("No menu or price information was found for this place.")

        sentiment = record.get("sentiment") or {}
        if sentiment.get("summary"):
            doc.add_heading("Review Sentiment", level=2)
            doc.add_paragraph(f"Overall: {sentiment.get('overall_sentiment')}")
            doc.add_paragraph(sentiment["summary"])

        doc.add_paragraph()  # spacing before the next place

    doc.styles["Normal"].font.size = Pt(11)

    doc.save(docx_path)
    return docx_path


def _flatten(record: dict, prefix: str = "") -> dict:
    """Flattens nested dicts into dotted keys for CSV, e.g. {"a": {"b": 1}}
    becomes {"a.b": 1}. Lists are stored as JSON text.
    """
    flat = {}
    for key, value in record.items():
        full_key = f"{prefix}{key}"
        if isinstance(value, dict):
            flat.update(_flatten(value, prefix=f"{full_key}."))
        elif isinstance(value, list):
            flat[full_key] = json.dumps(value, ensure_ascii=False)
        else:
            flat[full_key] = value
    return flat


# --- Product surveys ------------------------------------------------------------

_STATUS_LABELS = {
    "found_with_price": "Available, price found",
    "found_no_price": "Available, no price shown",
    "mentioned_in_reviews": "Likely available: mentioned in reviews",
    "no_evidence": "Needs field check",
    "inconclusive": "Couldn't check fully",
    "place_not_found": "Place not found on Google Maps",
    "menu_found": "Menu found",
    "menu_unavailable": "No menu found",
    "not_checked": "Skipped: the run was stopped first",
}

_SOURCE_LABELS = {
    "official_website": "Restaurant website",
    "third_party_search": "Online listing: delivery app, blog or social post (may be outdated)",
    "google_reviews": "Google review",
}


def save_survey(survey: dict, output_dir: str | os.PathLike, base_name: str, title: str) -> dict:
    """Writes a product survey or menu check as JSON (summary and places),
    CSV (one row per place) and DOCX. Returns the three paths; csv_path and
    docx_path are None when there are no places, and older files of the
    same name are then removed.
    """
    os.makedirs(output_dir, exist_ok=True)
    json_path = os.path.join(output_dir, f"{base_name}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(survey, f, ensure_ascii=False, indent=2)

    places = survey.get("places") or []
    csv_path = docx_path = None
    if places:
        csv_path = os.path.join(output_dir, f"{base_name}.csv")
        flat_rows = [_flatten(r) for r in places]
        fieldnames = sorted({key for row in flat_rows for key in row})
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(flat_rows)
        docx_path = save_survey_docx(survey, output_dir, base_name, title)
    else:
        _remove_stale(os.path.join(output_dir, f"{base_name}.csv"))
        _remove_stale(os.path.join(output_dir, f"{base_name}.docx"))
    return {"json_path": json_path, "csv_path": csv_path, "docx_path": docx_path}


def _add_table(doc, headers: list[str], rows: list[list]) -> None:
    """Adds a table with a header row and one row per entry of rows."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = "" if value is None else str(value)


def _fmt_number(value) -> str:
    """Formats a price for the report: whole numbers from 100 up, two decimals below."""
    return f"{value:,.0f}" if value >= 100 else f"{value:,.2f}"


def save_survey_docx(survey: dict, output_dir: str | os.PathLike, base_name: str, title: str) -> str | None:
    """Writes a readable survey report: what was checked, counts per status,
    a price range, an overview table, then each place's evidence or menu.
    Returns None when there are no places.
    """
    summary = survey.get("summary") or {}
    places = survey.get("places") or []
    if not places:
        return None
    is_product = summary.get("kind") == "product_survey"
    if is_product:
        places = [{**p, **verdict.explain(p)} for p in places]  # current wording, also for older files

    doc = Document()
    doc.add_heading(title, level=0)
    facts = doc.add_paragraph()
    for label, value in [
        ("Product", summary.get("product")),
        ("Also searched as", ", ".join(summary.get("aliases") or [])),
        ("Area", summary.get("area")),
        ("Place types", ", ".join(summary.get("place_types") or [])),
        ("Generated", summary.get("generated_at")),
    ]:
        if value:
            facts.add_run(f"{label}: ").bold = True
            facts.add_run(f"{value}\n")

    if is_product:
        doc.add_paragraph().add_run(
            "How to read this: every place has a one-line proof. \"Available\" means the "
            "product is listed on the place's own menu or an online listing such as a "
            "delivery app. \"Likely available\" means only a customer review mentions it. "
            "\"Needs field check\" means the product isn't mentioned online for that place "
            "yet, so a call or visit will confirm whether it's stocked."
        ).italic = True

    doc.add_heading("Summary", level=1)
    counts = summary.get("status_counts") or {}
    _add_table(doc, ["Result", "Places"], [[_STATUS_LABELS.get(k, k), v] for k, v in counts.items() if v])
    price = summary.get("price_summary")
    if price:
        doc.add_paragraph(
            f"Prices found: {price['count']} distinct price(s) in {price['currency'] or 'unknown currency'}, "
            f"from {_fmt_number(price['min'])} to {_fmt_number(price['max'])} "
            f"(median {_fmt_number(price['median'])}). Sizes and variants differ between "
            "places, so compare items, not just numbers."
        )

    doc.add_heading("Overview", level=1)
    rows = []
    for place in places:
        if is_product:
            best = verdict.best_menu_evidence(place)
            shown = best["price"] if best and best.get("price") else "-"
            rows.append([place.get("name") or place.get("searched_name"),
                         _STATUS_LABELS.get(place.get("status"), place.get("status")), shown, place.get("proof") or "-"])
        else:
            items = (place.get("menu") or {}).get("items") or []
            shown = f"{len(items)} item(s)" if items else "-"
            sources = _SOURCE_LABELS.get((place.get("menu") or {}).get("source"), "-")
            rows.append([place.get("name"), _STATUS_LABELS.get(place.get("status"), place.get("status")), shown, sources or "-"])
    _add_table(doc, ["Place", "Result", "Price", "Proof"] if is_product else ["Place", "Result", "Menu", "Source"], rows)

    doc.add_heading("Places", level=1)
    for place in places:
        doc.add_heading(place.get("name") or place.get("searched_name") or "Unknown", level=2)
        info = doc.add_paragraph()
        for label, key in [
            ("Searched for", "searched_name"),
            ("Address", "address"),
            ("Phone", "phone"),
            ("Google Maps", "google_maps_url"),
            ("Website", "website"),
        ]:
            if place.get(key) and not (key == "searched_name" and place.get(key) == place.get("name")):
                info.add_run(f"{label}: ").bold = True
                info.add_run(f"{place[key]}\n")
        info.add_run("Result: ").bold = True
        info.add_run(_STATUS_LABELS.get(place.get("status"), str(place.get("status"))))
        if place.get("proof"):
            info.add_run("\nProof: ").bold = True
            info.add_run(place["proof"])

        if is_product:
            evidence = place.get("evidence") or []
            if evidence:
                _add_table(
                    doc,
                    ["Item", "Price", "Source", "Quote"],
                    [[e.get("item"), e.get("price") or "-", _SOURCE_LABELS.get(e["source"], e["source"]), e.get("quote")] for e in evidence],
                )
            checks = place.get("checks") or {}
            if checks:
                doc.add_paragraph(
                    "Checked: " + "; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in checks.items())
                ).runs[0].italic = True
        else:
            menu = place.get("menu") or {}
            if menu.get("items"):
                _add_table(doc, ["Item", "Price"], [[i.get("item"), i.get("price")] for i in menu["items"]])
            if menu.get("note"):
                doc.add_paragraph(menu["note"]).runs[0].italic = True

        sources = place.get("third_party_sources") or (place.get("menu") or {}).get("sources") or []
        if sources:
            doc.add_paragraph("Search sources: " + ", ".join(s.get("title") or s.get("uri") or "" for s in sources))

    doc.styles["Normal"].font.size = Pt(10)
    docx_path = os.path.join(output_dir, f"{base_name}.docx")
    doc.save(docx_path)
    return docx_path
